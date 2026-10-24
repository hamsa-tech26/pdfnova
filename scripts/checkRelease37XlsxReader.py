#!/usr/bin/env python3
"""Independent parser check of an ephemeral, project-transcribed USGS workbook.

Reads the actual browser-downloaded XLSX with openpyxl. Input references are
created by the test from the existing project-curated publisher transcription.
This is NOT an independently reviewed gold reference or V4 Stable evidence.
Do not print PDF cell text, reference content or user file paths.
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

MAX_BYTES = 25 * 1024 * 1024
TABLES = 5
EXPECTED_CELLS = 147


def inspect(xlsx_path, reference_path):
    workbook_path = Path(xlsx_path)
    reference_file = Path(reference_path)
    if not workbook_path.is_file() or workbook_path.stat().st_size > MAX_BYTES:
        raise ValueError("Missing or oversized independent-read workbook")
    if not reference_file.is_file() or reference_file.stat().st_size > 100_000:
        raise ValueError("Missing or oversized source reference")
    reference = json.loads(reference_file.read_text(encoding="utf-8"))
    if (not isinstance(reference, list) or len(reference) != TABLES
            or not all(isinstance(rows, list) and rows and all(
                isinstance(row, list) and row and len(row) == len(rows[0])
                and all(isinstance(item, str) for item in row)
                for row in rows) for rows in reference)):
        raise ValueError("Invalid bounded project-transcribed reference")

    mismatch = 0
    formulas = 0
    observed = 0
    workbook = load_workbook(filename=workbook_path, read_only=True,
                             data_only=False, keep_links=True)
    try:
        if len(workbook.worksheets) != TABLES + 1:
            raise ValueError("Workbook sheet count differs from expected table and review sheets")
        if workbook._external_links:
            raise ValueError("Unexpected external workbook references")
        for index, rows in enumerate(reference):
            sheet = workbook.worksheets[index]
            width = len(rows[0])
            if sheet.max_row != len(rows) or sheet.max_column != width:
                raise ValueError("Workbook dimension differs from project-publisher table")
            for row_number, cells in enumerate(sheet.iter_rows(
                    min_row=1, max_row=len(rows), min_col=1,
                    max_col=width), start=0):
                for column_number, cell in enumerate(cells):
                    formulas += int(cell.data_type == "f")
                    value = "" if cell.value is None else str(cell.value)
                    if value != rows[row_number][column_number]:
                        mismatch += 1
                    observed += 1
        review_sheet = workbook.worksheets[-1]
        if "EXTRACTION NOT VERIFIED" not in str(review_sheet.cell(1, 1).value):
            raise ValueError("Workbook review-notes warning missing")
    finally:
        workbook.close()
    if observed != EXPECTED_CELLS or mismatch or formulas:
        raise ValueError("Independent workbook reader found mismatched cells or formulas")
    return {
        "status": "PASS", "reader": "openpyxl", "tables": TABLES,
        "comparedCells": observed, "mismatches": mismatch, "formulaCells": formulas,
        "referenceStatus": "PROJECT_TRANSCRIBED_NOT_INDEPENDENT",
        "formalV4Credit": 0,
    }


def main():
    if len(sys.argv) != 3:
        raise SystemExit("Usage: checkRelease37XlsxReader.py <xlsx> <private-reference>")
    try:
        result = inspect(sys.argv[1], sys.argv[2])
        print(json.dumps(result))
    except Exception:
        # Never leak source/reference cells or local file paths in CI logs.
        print(json.dumps({"status": "FAILED", "reason": "INDEPENDENT_READER_CHECK_FAILED",
                          "formalV4Credit": 0}))
        raise SystemExit(2)


if __name__ == "__main__":
    main()
